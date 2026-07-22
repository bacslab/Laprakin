import argparse
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ALIGNMENT = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def points(value):
    return round(value.pt, 2) if value else None


def centimeters(value):
    return round(value.cm, 2) if value else None


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def paragraph_summary(paragraph):
    text = clean(paragraph.text)
    runs = [run for run in paragraph.runs if clean(run.text)]
    fonts = Counter((run.font.name or "inherited") for run in runs)
    sizes = Counter(str(points(run.font.size) or "inherited") for run in runs)
    colors = Counter(
        str(run.font.color.rgb) if run.font.color and run.font.color.rgb else "inherited"
        for run in runs
    )
    return {
        "text": text[:220],
        "style": paragraph.style.name if paragraph.style else "",
        "alignment": ALIGNMENT.get(paragraph.alignment, "inherited"),
        "fonts": dict(fonts.most_common(3)),
        "sizes_pt": dict(sizes.most_common(3)),
        "colors": dict(colors.most_common(3)),
        "bold_runs": sum(1 for run in runs if run.bold),
        "italic_runs": sum(1 for run in runs if run.italic),
        "list": paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None,
    }


def is_heading(item):
    text = item["text"]
    style = item["style"].lower()
    if not text or len(text) > 180:
        return False
    if "heading" in style or "judul" in style or "title" in style:
        return True
    if re.match(r"^(bab\s+[ivxlcdm]+|\d+(?:\.\d+)*[.)]?\s+)", text, re.I):
        return True
    letters = [char for char in text if char.isalpha()]
    return len(letters) >= 4 and text == text.upper() and len(text.split()) <= 12


def analyze(path):
    document = Document(path)
    paragraphs = [paragraph_summary(paragraph) for paragraph in document.paragraphs]
    nonempty = [item for item in paragraphs if item["text"]]
    styles = Counter(item["style"] for item in nonempty)
    alignments = Counter(item["alignment"] for item in nonempty)
    fonts = Counter()
    sizes = Counter()
    colors = Counter()
    for item in nonempty:
        fonts.update(item["fonts"])
        sizes.update(item["sizes_pt"])
        colors.update(item["colors"])

    sections = []
    for section in document.sections:
        sections.append({
            "page_cm": [centimeters(section.page_width), centimeters(section.page_height)],
            "margins_cm": {
                "top": centimeters(section.top_margin),
                "right": centimeters(section.right_margin),
                "bottom": centimeters(section.bottom_margin),
                "left": centimeters(section.left_margin),
            },
            "header_cm": centimeters(section.header_distance),
            "footer_cm": centimeters(section.footer_distance),
            "header_text": clean(" ".join(p.text for p in section.header.paragraphs))[:180],
            "footer_text": clean(" ".join(p.text for p in section.footer.paragraphs))[:180],
        })

    tables = []
    for table in document.tables:
        preview = []
        for row in table.rows[:3]:
            preview.append([clean(cell.text)[:80] for cell in row.cells[:6]])
        tables.append({
            "rows": len(table.rows),
            "columns": len(table.columns),
            "style": table.style.name if table.style else "",
            "preview": preview,
        })

    image_count = sum(
        1 for relationship in document.part.rels.values()
        if "image" in relationship.reltype
    )
    return {
        "path": str(path),
        "paragraph_count": len(nonempty),
        "section_count": len(document.sections),
        "table_count": len(document.tables),
        "image_count": image_count,
        "style_counts": dict(styles.most_common(12)),
        "alignment_counts": dict(alignments.most_common()),
        "font_counts": dict(fonts.most_common(12)),
        "size_counts_pt": dict(sizes.most_common(12)),
        "color_counts": dict(colors.most_common(12)),
        "sections": sections,
        "candidate_headings": [item for item in nonempty if is_heading(item)][:80],
        "opening_paragraphs": nonempty[:35],
        "tables": tables,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    files = []
    for raw in args.inputs:
        candidate = Path(raw)
        if candidate.is_dir():
            files.extend(sorted(candidate.rglob("*.docx")))
        elif candidate.suffix.lower() == ".docx":
            files.append(candidate)

    result = []
    for path in files:
        try:
            result.append(analyze(path))
        except Exception as error:
            result.append({"path": str(path), "error": str(error)})

    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Distilled {len(result)} DOCX file(s) to {args.output}")


if __name__ == "__main__":
    main()
